"""
04_build_excel.py
Builds excel/revenue_budget_variance.xlsx — the Excel part of the project.
The budget is driven by editable growth targets (yellow cells on Assumptions);
every budget, variance and forecast figure is a live formula over the data sheets.
Afterwards run the LibreOffice recalc so calculated values are stored in the file.
"""
import sqlite3
import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = Path(__file__).resolve().parents[1]
OUT = BASE / "excel"; OUT.mkdir(exist_ok=True)
con = sqlite3.connect(BASE / "data" / "revenue.db")

# ------------------------------------------------------------------ styles
F, NAVY, BLUE, ORANGE, GREEN = "Arial", "1F3A5F", "2A78D6", "EB6834", "1BAF7A"
h_font, h_fill = Font(name=F, bold=True, color="FFFFFF", size=10), PatternFill("solid", fgColor=NAVY)
title_font, sub_font = Font(name=F, bold=True, size=16, color=NAVY), Font(name=F, italic=True, size=10, color="52514E")
bold, norm = Font(name=F, bold=True, size=10), Font(name=F, size=10)
blue_input = Font(name=F, bold=True, size=11, color="0000FF")
input_fill, kpi_fill = PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="EEF3FA")
fav_fill, unf_fill = PatternFill("solid", fgColor="DDF3EA"), PatternFill("solid", fgColor="FCE3D8")
thin = Side(style="thin", color="C9C8C2"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
MONEY, MONEY_M, PCT, PCT2, INT = '$#,##0;($#,##0);-', '$#,##0.0,,"M";($#,##0.0,,"M");-', '0.0%;(0.0%);-', '0.00%;(0.00%);-', '#,##0'
DATEF = "mmm yyyy"

def title(ws, t, s):
    ws["A1"], ws["A2"] = t, s; ws["A1"].font, ws["A2"].font = title_font, sub_font
    ws.sheet_view.showGridLines = False

def header(ws, row, col, labels, widths=None):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font, c.fill, c.border = h_font, h_fill, box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32
    for i, w in enumerate(widths or []):
        ws.column_dimensions[get_column_letter(col + i)].width = w

def put(ws, ref, value, fmt=None, font=norm, fill=None, align=None):
    c = ws[ref]; c.value = value; c.font = font; c.border = box
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if align: c.alignment = Alignment(horizontal=align)
    return c

def note(ws, ref, text):
    ws[ref] = text; ws[ref].font = sub_font

def labels_only_value():
    d = DataLabelList(); d.showVal, d.showSerName, d.showCatName, d.showLegendKey = True, False, False, False
    return d

def fav_rules(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], fill=fav_fill))
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], fill=unf_fill))

wb = Workbook()

# ================================================================== DATA SHEETS
def data_sheet(name, df, table_name, fmts=None, widths=None):
    ws = wb.create_sheet(name)
    header(ws, 1, 1, list(df.columns))
    for r, row in enumerate(df.itertuples(index=False), 2):
        for c, v in enumerate(row, 1):
            if pd.isna(v): continue
            cell = ws.cell(row=r, column=c, value=v.to_pydatetime() if isinstance(v, pd.Timestamp) else v)
            if fmts and df.columns[c - 1] in fmts: cell.number_format = fmts[df.columns[c - 1]]
    t = Table(displayName=table_name, ref=f"A1:{get_column_letter(df.shape[1])}{len(df) + 1}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True); ws.add_table(t)
    ws.freeze_panes = "A2"
    for i, col in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(col, max(12, len(col) + 3))
    return ws, len(df) + 1

store_month = pd.read_sql("SELECT store_id, store_type, size_sqft, month_start, fiscal_year, fiscal_month, "
                          "ROUND(revenue, 2) AS revenue FROM vw_monthly_store ORDER BY month_start, store_id",
                          con, parse_dates=["month_start"])
dept_month = pd.read_sql("""SELECT m.dept_id, s.store_type, m.fiscal_year, m.fiscal_month, ROUND(SUM(m.revenue), 2) AS revenue
                            FROM monthly_sales m JOIN stores s USING (store_id)
                            GROUP BY m.dept_id, s.store_type, m.fiscal_year, m.fiscal_month
                            ORDER BY m.dept_id, s.store_type, m.fiscal_year, m.fiscal_month""", con)
fc = pd.read_sql("SELECT * FROM forecast", con, parse_dates=["month_start"])
fc["store_type"] = fc.level_value.where(fc.level == "Store type")
fc["store_id"] = pd.to_numeric(fc.level_value.where(fc.level == "Store"), errors="coerce")
fc = fc[["level", "store_type", "store_id", "model", "purpose", "month_start", "forecast", "lower_80", "upper_80"]]
weekly = pd.read_sql("""SELECT week_end_date, ROUND(SUM(weekly_sales), 2) AS sales, MAX(is_holiday) AS is_holiday
                        FROM weekly_sales GROUP BY week_end_date ORDER BY week_end_date""", con, parse_dates=["week_end_date"])
stores = pd.read_sql("SELECT * FROM stores ORDER BY store_id", con)

M, m_last = data_sheet("Monthly_Data", store_month, "MonthlyStore",
                       {"month_start": "yyyy-mm-dd", "revenue": MONEY, "size_sqft": INT})
D, d_last = data_sheet("Dept_Data", dept_month, "DeptMonth", {"revenue": MONEY})
FC, f_last = data_sheet("Forecast_Data", fc, "Forecasts",
                        {"month_start": "yyyy-mm-dd", "forecast": MONEY, "lower_80": MONEY, "upper_80": MONEY},
                        {"model": 26, "level": 12})
W, w_last = data_sheet("Weekly_Total", weekly, "WeeklyTotal", {"week_end_date": "yyyy-mm-dd", "sales": MONEY})

def rng(sheet, col, last):  # absolute column range on a data sheet
    return f"{sheet}!${col}$2:${col}${last}"
MS = {k: rng("Monthly_Data", c, m_last) for k, c in zip(["store", "type", "size", "month", "fy", "fm", "rev"], "ABCDEFG")}
DD = {k: rng("Dept_Data", c, d_last) for k, c in zip(["dept", "type", "fy", "fm", "rev"], "ABCDE")}
FF = {k: rng("Forecast_Data", c, f_last) for k, c in zip(["level", "type", "store", "model", "purpose", "month", "fc", "lo", "hi"], "ABCDEFGHI")}
WW = {k: rng("Weekly_Total", c, w_last) for k, c in zip(["week", "sales", "hol"], "ABC")}

# ================================================================== ASSUMPTIONS
A = wb.create_sheet("Assumptions", 0)
title(A, "Budget Assumptions", "Change the yellow cells - every budget and variance in the workbook updates.")
header(A, 4, 2, ["Store type", "FY2013 growth target", "Rationale"], [12, 18, 46])
for i, (t, g, why) in enumerate([("A", 0.03, "Large supercenters - main growth driver"),
                                 ("B", 0.02, "Mid-size stores - modest growth"),
                                 ("C", 0.015, "Small stores - low growth")]):
    r = 5 + i
    put(A, f"B{r}", t, font=bold, align="center")
    put(A, f"C{r}", g, "0.0%", font=blue_input, fill=input_fill, align="center")
    put(A, f"D{r}", why)
put(A, "B9", "Adjust budget for trading days?", font=bold); A.merge_cells("B9:C9")
put(A, "D9", "No", font=blue_input, fill=input_fill, align="center")
dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False); A.add_data_validation(dv); dv.add("D9")
note(A, "B10", "No  = budget is last year x (1 + growth) - the method used in SQL and Python.")
note(A, "B11", "Yes = also scales each month by its number of days (fixes the Feb 2012 leap day: 29 vs 28 days).")
put(A, "B13", "Budget method", font=bold)
note(A, "B14", "FY2013 budget (Feb 2012 - Jan 2013) = FY2012 actual for the same store / month x (1 + growth target for the store type).")
note(A, "B15", "These targets are ASSUMPTIONS for the project - Walmart's internal budget is not public. All sales figures are real.")
GROWTH = {"A": "Assumptions!$C$5", "B": "Assumptions!$C$6", "C": "Assumptions!$C$7"}
TOGGLE = "Assumptions!$D$9"

def day_factor(month_ref):
    return f'IF({TOGGLE}="Yes",DAY(EOMONTH({month_ref},0))/DAY(EOMONTH(EDATE({month_ref},-12),0)),1)'

# ================================================================== BUDGET vs ACTUAL (monthly)
B = wb.create_sheet("Budget_vs_Actual", 1)
title(B, "FY2013 Budget vs Actual, by Month", "Actuals Feb - Sep 2012; Oct 2012 - Jan 2013 use the Python Holt-Winters forecast. Positive variance = favourable.")
cols = ["Month", "Fiscal month", "LY actual - Type A", "LY actual - Type B", "LY actual - Type C", "Day factor",
        "Budget", "Actual", "Variance", "Variance %", "Status", "Forecast", "Latest estimate", "LE vs budget", "Cumulative LE vs budget"]
header(B, 4, 1, cols, [11, 8, 15, 15, 15, 8, 15, 15, 14, 10, 13, 15, 16, 14, 16])
months = pd.date_range("2012-02-01", "2013-01-01", freq="MS")
for i, mth in enumerate(months):
    r = 5 + i
    put(B, f"A{r}", mth.to_pydatetime(), DATEF, bold)
    put(B, f"B{r}", i + 1, "0", align="center")
    for j, t in enumerate("ABC"):
        put(B, f"{'CDE'[j]}{r}", f'=SUMIFS({MS["rev"]},{MS["fy"]},2012,{MS["fm"]},$B{r},{MS["type"]},"{t}")', MONEY)
    put(B, f"F{r}", f"={day_factor(f'$A{r}')}", "0.000")
    put(B, f"G{r}", f'=(C{r}*(1+{GROWTH["A"]})+D{r}*(1+{GROWTH["B"]})+E{r}*(1+{GROWTH["C"]}))*F{r}', MONEY)
    put(B, f"H{r}", f'=IF(COUNTIFS({MS["fy"]},2013,{MS["fm"]},$B{r})=0,"",SUMIFS({MS["rev"]},{MS["fy"]},2013,{MS["fm"]},$B{r}))', MONEY)
    put(B, f"I{r}", f'=IF(H{r}="","",H{r}-G{r})', MONEY)
    put(B, f"J{r}", f'=IF(H{r}="","",I{r}/G{r})', PCT2)
    put(B, f"K{r}", f'=IF(H{r}="","Forecast",IF(I{r}>=0,"Favourable","Unfavourable"))', align="center")
    put(B, f"L{r}", f'=IF(H{r}<>"","",SUMIFS({FF["fc"]},{FF["level"]},"Total",{FF["purpose"]},"Forecast",{FF["month"]},$A{r}))', MONEY)
    put(B, f"M{r}", f'=IF(H{r}<>"",H{r},L{r})', MONEY)
    put(B, f"N{r}", f"=M{r}-G{r}", MONEY)
    put(B, f"O{r}", f"=SUM($N$5:N{r})", MONEY)
tr = 17
put(B, f"A{tr}", "Full year", font=bold)
for c in "CDEGLMNH":
    put(B, f"{c}{tr}", f"=SUM({c}5:{c}16)", MONEY, bold)
put(B, f"I{tr}", f"=SUM(I5:I16)", MONEY, bold)
put(B, f"N{tr}", f"=M{tr}-G{tr}", MONEY, bold)
put(B, f"J{tr}", f"=IFERROR(N{tr}/G{tr},0)", PCT2, bold)
note(B, f"J{tr + 1}", "J17 = full-year latest estimate vs budget")
yr = 19
put(B, f"A{yr}", "Year to date (months with actuals)", font=bold); B.merge_cells(f"A{yr}:F{yr}")
put(B, f"G{yr}", f'=SUMPRODUCT((H5:H16<>"")*G5:G16)', MONEY, bold)
put(B, f"H{yr}", "=SUM(H5:H16)", MONEY, bold)
put(B, f"I{yr}", f"=H{yr}-G{yr}", MONEY, bold)
put(B, f"J{yr}", f"=IFERROR(I{yr}/G{yr},0)", PCT2, bold)
fav_rules(B, "I5:I16"); fav_rules(B, "N5:N17"); fav_rules(B, f"I{yr}:J{yr}")
B.conditional_formatting.add("K5:K16", CellIsRule(operator="equal", formula=['"Favourable"'], fill=fav_fill))
B.conditional_formatting.add("K5:K16", CellIsRule(operator="equal", formula=['"Unfavourable"'], fill=unf_fill))
B.freeze_panes = "B5"

bar = BarChart(); bar.type = "col"; bar.title = "Latest estimate (actual + forecast) vs budget"; bar.style = 10
bar.add_data(Reference(B, min_col=13, min_row=4, max_row=16), titles_from_data=True)
bar.set_categories(Reference(B, min_col=1, min_row=5, max_row=16))
bar.series[0].graphicalProperties.solidFill = BLUE
line = LineChart(); line.add_data(Reference(B, min_col=7, min_row=4, max_row=16), titles_from_data=True)
line.series[0].graphicalProperties.line.solidFill = "52514E"; line.series[0].graphicalProperties.line.width = 28000
bar += line
bar.y_axis.numFmt = '$#,##0,,"M"'; bar.height, bar.width = 8.5, 22; bar.legend.position = "b"
bar.y_axis.scaling.min = 150_000_000
B.add_chart(bar, "A22")
vc = BarChart(); vc.type = "col"; vc.title = "Monthly variance vs budget (actual months)"; vc.style = 10
vc.add_data(Reference(B, min_col=9, min_row=4, max_row=12), titles_from_data=True)
vc.set_categories(Reference(B, min_col=1, min_row=5, max_row=12)); vc.legend = None
vc.series[0].graphicalProperties.solidFill = BLUE; vc.y_axis.numFmt = '$#,##0.0,,"M"'; vc.height, vc.width = 8.5, 16
B.add_chart(vc, "K22")

# ================================================================== STORE BvA
S = wb.create_sheet("Store_BvA", 2)
title(S, "Budget vs Actual by Store", "Year to date = Feb - Sep 2012. Full-year forecast for Oct - Jan uses each store's run-rate (Python).")
scols = ["Store", "Type", "Size (sq ft)", "LY YTD (Feb-Sep 2011)", "LY Feb 2011", "YTD budget", "YTD actual", "YTD variance",
         "Variance %", "Rank (1 = best)", "Status", "FY2012 revenue", "Revenue / sq ft", "Full-year budget",
         "Forecast Oct-Jan", "Latest estimate", "LE vs budget", "LE vs budget %"]
header(S, 4, 1, scols, [7, 6, 11, 16, 13, 15, 15, 13, 10, 9, 13, 15, 10, 15, 15, 15, 13, 10])
n = len(stores)
for i, st in enumerate(stores.itertuples()):
    r = 5 + i
    g = f'IF(B{r}="A",{GROWTH["A"]},IF(B{r}="B",{GROWTH["B"]},{GROWTH["C"]}))'
    feb_f = day_factor("DATE(2012,2,1)")
    put(S, f"A{r}", int(st.store_id), "0", bold, align="center")
    put(S, f"B{r}", st.store_type, align="center")
    put(S, f"C{r}", int(st.size_sqft), INT)
    put(S, f"D{r}", f'=SUMIFS({MS["rev"]},{MS["store"]},$A{r},{MS["fy"]},2012,{MS["fm"]},"<=8")', MONEY)
    put(S, f"E{r}", f'=SUMIFS({MS["rev"]},{MS["store"]},$A{r},{MS["fy"]},2012,{MS["fm"]},1)', MONEY)
    put(S, f"F{r}", f"=(D{r}+E{r}*({feb_f}-1))*(1+{g})", MONEY)
    put(S, f"G{r}", f'=SUMIFS({MS["rev"]},{MS["store"]},$A{r},{MS["fy"]},2013)', MONEY)
    put(S, f"H{r}", f"=G{r}-F{r}", MONEY)
    put(S, f"I{r}", f"=IFERROR(H{r}/F{r},0)", PCT)
    put(S, f"J{r}", f"=RANK(I{r},$I$5:$I${4 + n})", "0", align="center")
    put(S, f"K{r}", f'=IF(H{r}>=0,"On / above","Below")', align="center")
    put(S, f"L{r}", f'=SUMIFS({MS["rev"]},{MS["store"]},$A{r},{MS["fy"]},2012)', MONEY)
    put(S, f"M{r}", f"=IFERROR(L{r}/C{r},0)", "$#,##0")
    put(S, f"N{r}", f"=(L{r}+E{r}*({feb_f}-1))*(1+{g})", MONEY)
    put(S, f"O{r}", f'=SUMIFS({FF["fc"]},{FF["level"]},"Store",{FF["store"]},$A{r})', MONEY)
    put(S, f"P{r}", f"=G{r}+O{r}", MONEY)
    put(S, f"Q{r}", f"=P{r}-N{r}", MONEY)
    put(S, f"R{r}", f"=IFERROR(Q{r}/N{r},0)", PCT)
last = 4 + n; tot = last + 1
put(S, f"A{tot}", "Total", font=bold)
for c in "CDEFGHLNOPQ":
    put(S, f"{c}{tot}", f"=SUM({c}5:{c}{last})", MONEY if c != "C" else INT, bold)
put(S, f"I{tot}", f"=IFERROR(H{tot}/F{tot},0)", PCT2, bold); put(S, f"R{tot}", f"=IFERROR(Q{tot}/N{tot},0)", PCT2, bold)
put(S, f"K{tot}", f'=COUNTIF(K5:K{last},"On / above")&" of {n} on/above"', font=bold)
fav_rules(S, f"H5:I{last}"); fav_rules(S, f"Q5:R{last}")
S.conditional_formatting.add(f"M5:M{last}", DataBarRule(start_type="min", end_type="max", color=BLUE))
S.freeze_panes = "C5"; S.auto_filter.ref = f"A4:R{last}"
note(S, f"A{tot + 2}", "Tip: use the filter arrows on row 4 to sort stores by variance %. Store 14 is the largest $ shortfall.")

# store-type summary under the store table
tt = tot + 4
put(S, f"A{tt - 1}", "Summary by store type", font=Font(name=F, bold=True, size=12, color=NAVY))
header(S, tt, 1, ["Type", "Stores", "Growth target", "YTD budget", "YTD actual", "YTD variance", "Variance %",
                  "Full-year budget", "Latest estimate", "LE vs budget %"])
for i, t in enumerate("ABC"):
    r = tt + 1 + i
    put(S, f"A{r}", t, font=bold, align="center")
    put(S, f"B{r}", f'=COUNTIF($B$5:$B${last},A{r})', "0", align="center")
    put(S, f"C{r}", f"={GROWTH[t]}", "0.0%", align="center")
    for c, src in zip("DEF", "FGH"):
        put(S, f"{c}{r}", f"=SUMIFS(${src}$5:${src}${last},$B$5:$B${last},$A{r})", MONEY)
    put(S, f"G{r}", f"=IFERROR(F{r}/D{r},0)", PCT2)
    put(S, f"H{r}", f"=SUMIFS($N$5:$N${last},$B$5:$B${last},$A{r})", MONEY)
    put(S, f"I{r}", f"=SUMIFS($P$5:$P${last},$B$5:$B${last},$A{r})", MONEY)
    put(S, f"J{r}", f"=IFERROR(I{r}/H{r}-1,0)", PCT2)
fav_rules(S, f"F{tt + 1}:G{tt + 3}"); fav_rules(S, f"J{tt + 1}:J{tt + 3}")
ch = BarChart(); ch.type = "col"; ch.title = "YTD variance % by store type"; ch.style = 10
ch.add_data(Reference(S, min_col=7, min_row=tt, max_row=tt + 3), titles_from_data=True)
ch.set_categories(Reference(S, min_col=1, min_row=tt + 1, max_row=tt + 3)); ch.legend = None
ch.series[0].graphicalProperties.solidFill = BLUE; ch.y_axis.numFmt = "0.0%"; ch.dataLabels = labels_only_value()
ch.height, ch.width = 7, 12
S.add_chart(ch, f"L{tt - 1}")

# ================================================================== DEPARTMENT BvA
DP = wb.create_sheet("Dept_BvA", 3)
title(DP, "Budget vs Actual by Department (YTD)", "Department budgets follow the same rule: LY actual by store type x (1 + growth target).")
header(DP, 4, 1, ["Dept", "LY YTD - A", "LY YTD - B", "LY YTD - C", "YTD budget", "YTD actual", "Variance", "Variance %",
                  "Abs. variance rank"], [8, 14, 14, 14, 15, 15, 13, 10, 10])
depts = sorted(dept_month.dept_id.unique())
for i, dpt in enumerate(depts):
    r = 5 + i
    put(DP, f"A{r}", int(dpt), "0", bold, align="center")
    for j, t in enumerate("ABC"):
        put(DP, f"{'BCD'[j]}{r}", f'=SUMIFS({DD["rev"]},{DD["dept"]},$A{r},{DD["fy"]},2012,{DD["fm"]},"<=8",{DD["type"]},"{t}")', MONEY)
    put(DP, f"E{r}", f'=B{r}*(1+{GROWTH["A"]})+C{r}*(1+{GROWTH["B"]})+D{r}*(1+{GROWTH["C"]})', MONEY)
    put(DP, f"F{r}", f'=SUMIFS({DD["rev"]},{DD["dept"]},$A{r},{DD["fy"]},2013)', MONEY)
    put(DP, f"G{r}", f"=F{r}-E{r}", MONEY)
    put(DP, f"H{r}", f'=IF(E{r}<=0,"",G{r}/E{r})', PCT)
    put(DP, f"I{r}", f"=RANK(ABS(G{r}),$K$5:$K${4 + len(depts)})", "0", align="center")
    DP[f"K{r}"] = f"=ABS(G{r})"; DP[f"K{r}"].number_format = MONEY; DP[f"K{r}"].font = Font(name=F, size=8, color="999999")
dl = 4 + len(depts)
put(DP, f"A{dl + 1}", "Total", font=bold)
for c in "BCDEFG": put(DP, f"{c}{dl + 1}", f"=SUM({c}5:{c}{dl})", MONEY, bold)
fav_rules(DP, f"G5:H{dl}")
DP["K4"] = "helper: |variance|"; DP["K4"].font = Font(name=F, size=8, color="999999")
note(DP, "M4", "Tip: filter or sort column I to see the departments that move the total most.")
note(DP, "M5", "Note: budget here does not apply the trading-day factor (department detail is kept simple).")
DP.freeze_panes = "B5"; DP.auto_filter.ref = f"A4:I{dl}"

# ================================================================== EXCEL FORECAST (seasonal index method)
X = wb.create_sheet("Excel_Forecast", 4)
title(X, "Forecast in Excel: Seasonal Index x Trend", "Classic FP&A method built only with formulas. Compare its accuracy with the budget and Python's Holt-Winters.")
header(X, 4, 1, ["Month", "t", "FY", "Fiscal month", "Revenue", "Days", "Revenue / day", "Seasonal index",
                 "Deseasonalised / day", "Trend (train to Jan 2012)", "Backtest forecast", "Abs % error"],
       [11, 5, 7, 8, 15, 6, 13, 10, 14, 14, 15, 10])
all_months = pd.date_range("2010-02-01", "2013-01-01", freq="MS")
first, n_hist = 5, 32
for i, mth in enumerate(all_months):
    r = first + i
    fy_ = mth.year + (0 if mth.month == 1 else 1); fm_ = (mth.month - 2) % 12 + 1
    put(X, f"A{r}", mth.to_pydatetime(), DATEF, bold)
    put(X, f"B{r}", i + 1, "0", align="center"); put(X, f"C{r}", fy_, "0", align="center"); put(X, f"D{r}", fm_, "0", align="center")
    put(X, f"F{r}", f"=DAY(EOMONTH(A{r},0))", "0", align="center")
    put(X, f"H{r}", f"=INDEX($P$5:$P$16,D{r})", "0.000")
    if i < n_hist:
        put(X, f"E{r}", f'=SUMIFS({MS["rev"]},{MS["fy"]},C{r},{MS["fm"]},D{r})', MONEY)
        put(X, f"G{r}", f"=E{r}/F{r}", MONEY)
        put(X, f"I{r}", f"=G{r}/H{r}", MONEY)
    put(X, f"J{r}", f"=FORECAST(B{r},$I$5:$I$28,$B$5:$B$28)", MONEY)
    if 24 <= i < n_hist:
        put(X, f"K{r}", f"=J{r}*H{r}*F{r}", MONEY)
        put(X, f"L{r}", f"=ABS(K{r}-E{r})/E{r}", PCT2)
# seasonal index table
header(X, 4, 14, ["Fiscal month", "Month", "Seasonal index"], [8, 7, 10])
for k in range(12):
    r = 5 + k
    put(X, f"N{r}", k + 1, "0", align="center")
    put(X, f"O{r}", ["Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan"][k], align="center")
    put(X, f"P{r}", (f"=AVERAGE(INDEX($G$5:$G$16,N{r})/AVERAGE($G$5:$G$16),"
                     f"INDEX($G$17:$G$28,N{r})/AVERAGE($G$17:$G$28))"), "0.000")
X.conditional_formatting.add("P5:P16", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color=ORANGE))
note(X, "N18", "Index = month's revenue/day / that year's average, averaged over FY2011 & FY2012.")

# accuracy comparison
put(X, "N21", "Backtest accuracy, Feb - Sep 2012 (MAPE = average % miss)", font=Font(name=F, bold=True, size=11, color=NAVY))
header(X, 22, 14, ["Method", "", "MAPE"])
X.column_dimensions["N"].width = 30
put(X, "N23", "Excel: seasonal index x trend", font=bold); put(X, "P23", f"=AVERAGE(L{first + 24}:L{first + 31})", PCT2)
put(X, "N24", "Budget: last year x growth", font=bold)
put(X, "P24", "=SUMPRODUCT(ABS(Budget_vs_Actual!G5:G12-Budget_vs_Actual!H5:H12)/Budget_vs_Actual!H5:H12)/8", PCT2)
put(X, "N25", "Python: Holt-Winters", font=bold)
hw_parts = "+".join([f'ABS(SUMIFS({FF["fc"]},{FF["model"]},"Holt-Winters",{FF["purpose"]},"Backtest",{FF["month"]},A{first + 24 + k})-E{first + 24 + k})/E{first + 24 + k}' for k in range(8)])
put(X, "P25", f"=({hw_parts})/8", PCT2)
X.conditional_formatting.add("P23:P25", ColorScaleRule(start_type="min", start_color=GREEN, end_type="max", end_color="FFFFFF"))

# forward forecast with the Excel method (refit trend on all 32 months)
put(X, "N28", "Excel forecast for Oct 2012 - Jan 2013 (trend refit on all 32 months)", font=Font(name=F, bold=True, size=11, color=NAVY))
header(X, 29, 14, ["Month", "Excel forecast", "Python Holt-Winters", "Budget"])
X.column_dimensions["Q"].width = 15
for k in range(4):
    r = 30 + k; src = first + 32 + k
    put(X, f"N{r}", f"=A{src}", DATEF, bold)
    put(X, f"O{r}", f"=FORECAST(B{src},$I$5:$I$36,$B$5:$B$36)*H{src}*F{src}", MONEY)
    put(X, f"P{r}", f'=SUMIFS({FF["fc"]},{FF["level"]},"Total",{FF["purpose"]},"Forecast",{FF["month"]},N{r})', MONEY)
    put(X, f"Q{r}", f"=INDEX(Budget_vs_Actual!$G$5:$G$16,MATCH(N{r},Budget_vs_Actual!$A$5:$A$16,0))", MONEY)
put(X, "N34", "Total Oct - Jan", font=bold)
for c in "OPQ": put(X, f"{c}34", f"=SUM({c}30:{c}33)", MONEY, bold)
X.column_dimensions["O"].width = 15; X.column_dimensions["P"].width = 16
lc = LineChart(); lc.title = "Revenue: actual and Excel seasonal-index backtest"; lc.style = 12
lc.add_data(Reference(X, min_col=5, min_row=4, max_row=36), titles_from_data=True)
lc.add_data(Reference(X, min_col=11, min_row=4, max_row=36), titles_from_data=True)
lc.set_categories(Reference(X, min_col=1, min_row=5, max_row=36))
lc.series[0].graphicalProperties.line.solidFill = "52514E"; lc.series[1].graphicalProperties.line.solidFill = ORANGE
lc.series[1].graphicalProperties.line.dashStyle = "dash"; lc.y_axis.numFmt = '$#,##0,,"M"'
lc.y_axis.scaling.min = 170_000_000; lc.height, lc.width = 8, 20; lc.legend.position = "b"
X.add_chart(lc, "A43")

# ================================================================== HOLIDAYS
HL = wb.create_sheet("Holidays", 5)
title(HL, "Holiday Weeks and Peak Weeks", "Company total weekly sales (Sat - Fri weeks).")
header(HL, 4, 1, ["Week type", "Weeks", "Avg weekly sales", "vs normal week"], [24, 9, 17, 12])
put(HL, "A5", "Normal week", font=bold); put(HL, "A6", "Flagged holiday week", font=bold)
for r, flag in ((5, 0), (6, 1)):
    put(HL, f"B{r}", f"=COUNTIF({WW['hol']},{flag})", "0", align="center")
    put(HL, f"C{r}", f"=AVERAGEIF({WW['hol']},{flag},{WW['sales']})", MONEY)
    put(HL, f"D{r}", f"=C{r}/$C$5-1", PCT)
header(HL, 9, 1, ["Rank", "Week ending", "Sales", "Holiday flag?"])
for k in range(1, 6):
    r = 9 + k
    put(HL, f"A{r}", k, "0", align="center")
    put(HL, f"C{r}", f"=LARGE({WW['sales']},A{r})", MONEY)
    put(HL, f"B{r}", f"=INDEX({WW['week']},MATCH(C{r},{WW['sales']},0))", "yyyy-mm-dd")
    put(HL, f"D{r}", f'=IF(INDEX({WW["hol"]},MATCH(C{r},{WW["sales"]},0))=1,"Yes","No")', align="center")
note(HL, "A16", "The two biggest weeks are the weeks BEFORE Christmas, which Walmart's holiday flag does not mark.")
lc = LineChart(); lc.title = "Weekly sales, all stores"; lc.style = 12
lc.add_data(Reference(W, min_col=2, min_row=1, max_row=w_last), titles_from_data=True)
lc.set_categories(Reference(W, min_col=1, min_row=2, max_row=w_last)); lc.legend = None
lc.series[0].graphicalProperties.line.solidFill = BLUE; lc.y_axis.numFmt = '$#,##0,,"M"'
lc.x_axis.number_format = "mmm yy"; lc.height, lc.width = 8, 22
HL.add_chart(lc, "F4")

# ================================================================== CLEANING LOG
CL = wb.create_sheet("Cleaning_Log")
title(CL, "Data Checks and Cleaning", "From scripts/01_prepare_data.py - the original files in data/source were never modified.")
log = pd.read_csv(BASE / "data" / "clean" / "cleaning_log.csv")
header(CL, 4, 1, list(log.columns), [20, 48, 10, 80])
for i, row in enumerate(log.itertuples(index=False)):
    for j, v in enumerate(row):
        c = CL.cell(row=5 + i, column=j + 1, value=v); c.font = norm; c.border = box
        if j == 2: c.number_format = INT

# ================================================================== DASHBOARD
DB = wb.create_sheet("Dashboard", 0)
title(DB, "Walmart FY2013 Revenue: Budget vs Actual & Forecast", "Real Walmart sales, 45 stores | FY2013 = Feb 2012 - Jan 2013 | budget = assumption (see Assumptions sheet)")
for c, w in zip("ABCDEFGHIJK", [2, 32, 15, 2, 33, 15, 2, 32, 15, 2, 10]):
    DB.column_dimensions[c].width = w
tiles = [
    ("B4", "C4", "FY2012 revenue (last full year)", f'=SUMIFS({MS["rev"]},{MS["fy"]},2012)', MONEY_M),
    ("E4", "F4", "FY2013 YTD actual (Feb-Sep)", "=Budget_vs_Actual!H19", MONEY_M),
    ("H4", "I4", "YTD growth vs last year", f'=Budget_vs_Actual!H19/SUMIFS({MS["rev"]},{MS["fy"]},2012,{MS["fm"]},"<=8")-1', PCT2),
    ("B6", "C6", "FY2013 YTD budget", "=Budget_vs_Actual!G19", MONEY_M),
    ("E6", "F6", "YTD variance vs budget", "=Budget_vs_Actual!I19", MONEY_M),
    ("H6", "I6", "YTD variance %", "=Budget_vs_Actual!J19", PCT2),
    ("B8", "C8", "Full-year budget", "=Budget_vs_Actual!G17", MONEY_M),
    ("E8", "F8", "Latest estimate (actual + forecast)", "=Budget_vs_Actual!M17", MONEY_M),
    ("H8", "I8", "Latest estimate vs budget", "=Budget_vs_Actual!J17", PCT2),
    ("B10", "C10", "Stores on / above budget (YTD)", f'=COUNTIF(Store_BvA!K5:K{last},"On / above")', "0"),
    ("E10", "F10", "Leap-day sales in Feb 2012", "=Budget_vs_Actual!H5/29", MONEY_M),
    ("H10", "I10", "Best backtest MAPE (Holt-Winters)", "=Excel_Forecast!P25", PCT2),
]
for lr, vr, lab, f, fmt in tiles:
    put(DB, lr, lab, font=bold, fill=kpi_fill)
    c = put(DB, vr, f, fmt, Font(name=F, bold=True, size=14, color=NAVY), kpi_fill); c.alignment = Alignment(horizontal="right")
for r in (4, 6, 8, 10): DB.row_dimensions[r].height = 24
fav_rules(DB, "F6"); fav_rules(DB, "I6"); fav_rules(DB, "I8")
DB["B12"] = "Key findings"; DB["B12"].font = Font(name=F, bold=True, size=12, color=NAVY)
for i, t in enumerate([
    "1. YTD revenue is slightly ahead of budget, but most of February's beat is the extra leap day (switch the trading-day option on the Assumptions sheet to see it).",
    "2. Momentum slowed: Feb-May beat budget, Jun-Aug fell behind.",
    "3. Type A stores (3% target) are just behind budget; Types B and C beat lower targets.",
    "4. Store 14 alone is about $7.7M behind - almost as much as the company's whole YTD lead.",
    "5. The forecast puts the full year within about 1% of budget; October and November carry the most risk.",
]):
    DB[f"B{13 + i}"] = t; DB[f"B{13 + i}"].font = norm
note(DB, "B19", "Sheets: Assumptions (change growth targets) | Budget_vs_Actual | Store_BvA | Dept_BvA | Excel_Forecast | Holidays | data sheets | Cleaning_Log")
ch = BarChart(); ch.type = "col"; ch.title = "Latest estimate vs budget by month"; ch.style = 10
ch.add_data(Reference(B, min_col=13, min_row=4, max_row=16), titles_from_data=True)
ch.set_categories(Reference(B, min_col=1, min_row=5, max_row=16))
ch.series[0].graphicalProperties.solidFill = BLUE
ln = LineChart(); ln.add_data(Reference(B, min_col=7, min_row=4, max_row=16), titles_from_data=True)
ln.series[0].graphicalProperties.line.solidFill = "52514E"; ln.series[0].graphicalProperties.line.width = 28000
ch += ln; ch.y_axis.numFmt = '$#,##0,,"M"'; ch.y_axis.scaling.min = 150_000_000
ch.height, ch.width = 8.5, 17; ch.legend.position = "b"
DB.add_chart(ch, "B21")
ch2 = BarChart(); ch2.type = "col"; ch2.title = "YTD variance % by store type"; ch2.style = 10
ch2.add_data(Reference(S, min_col=7, min_row=tt, max_row=tt + 3), titles_from_data=True)
ch2.set_categories(Reference(S, min_col=1, min_row=tt + 1, max_row=tt + 3)); ch2.legend = None
ch2.series[0].graphicalProperties.solidFill = BLUE; ch2.y_axis.numFmt = "0.0%"; ch2.dataLabels = labels_only_value()
ch2.height, ch2.width = 8.5, 11
DB.add_chart(ch2, "H21")

# ================================================================== README
R = wb.create_sheet("README", 0)
title(R, "Monthly Revenue Forecasting & Budget Variance - Excel workbook",
      "Real Walmart weekly sales (45 stores, Feb 2010 - Oct 2012), allocated to months in SQL. Budget = assumption.")
R.column_dimensions["A"].width = 22; R.column_dimensions["B"].width = 110
rows = [("Sheet", "What it does"),
        ("Dashboard", "Headline KPIs, key findings and two charts."),
        ("Assumptions", "EDIT HERE: growth targets by store type and the trading-day option. Everything else recalculates."),
        ("Budget_vs_Actual", "FY2013 by month: last-year actual, budget, actual, variance, status, forecast and latest estimate."),
        ("Store_BvA", "Every store: YTD budget vs actual, rank, revenue per sq ft, full-year latest estimate; summary by store type."),
        ("Dept_BvA", "81 departments: YTD budget vs actual and which departments move the total most."),
        ("Excel_Forecast", "Seasonal index x trend forecast built with FORECAST/INDEX formulas; backtest MAPE vs budget and Python."),
        ("Holidays", "Holiday vs normal weeks and the five biggest weeks (LARGE + INDEX/MATCH)."),
        ("Monthly_Data", "Store x month revenue (Feb 2010 - Sep 2012), from the SQL view vw_monthly_store."),
        ("Dept_Data", "Department x store type x month revenue."),
        ("Forecast_Data", "Forecasts written by the Python notebook (backtest and Oct 2012 - Jan 2013)."),
        ("Weekly_Total", "Company total sales per week with Walmart's holiday flag."),
        ("Cleaning_Log", "Data checks and cleaning applied to the original files."),
        ("", ""),
        ("Formulas used", "SUMIFS, COUNTIFS, AVERAGEIF, INDEX/MATCH, RANK, LARGE, FORECAST, EOMONTH, EDATE, SUMPRODUCT, IF/IFERROR."),
        ("Colour key", "Yellow + blue text = inputs. Green fill = favourable variance, orange fill = unfavourable."),
        ("Sign convention", "Variance = actual - budget. Positive = favourable (revenue above budget).")]
for i, (a, b) in enumerate(rows):
    r = 4 + i; R[f"A{r}"], R[f"B{r}"] = a, b
    R[f"A{r}"].font = h_font if i == 0 else bold; R[f"B{r}"].font = h_font if i == 0 else norm
    if i == 0: R[f"A{r}"].fill = R[f"B{r}"].fill = h_fill

R.sheet_properties.tabColor = "52514E"; DB.sheet_properties.tabColor = NAVY; A.sheet_properties.tabColor = "EDA100"
# straight (not smoothed) lines in every line chart
for ws_ in wb.worksheets:
    for chart in ws_._charts:
        for sub in [chart] + list(getattr(chart, "_charts", [])):
            for ser in getattr(sub, "series", []):
                ser.smooth = False
wb.remove(wb["Sheet"])
wb.active = 1
path = OUT / "revenue_budget_variance.xlsx"
wb.save(path); print("saved", path)
